from __future__ import annotations

from pathlib import Path


def recognize_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            lines.append(f"## Лист: {name}")
            for row in ws.iter_rows(values_only=True):
                lines.append(
                    "\t".join("" if c is None else str(c) for c in row),
                )
            lines.append("")
    finally:
        wb.close()
    return "\n".join(lines).strip()
