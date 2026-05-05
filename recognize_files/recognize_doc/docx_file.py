from __future__ import annotations

from pathlib import Path

from docx.table import Table
from docx.text.paragraph import Paragraph


def _norm(s: str) -> str:
    return " ".join((s or "").replace("\n", " ").split())


def _cell_full_text(cell) -> str:
    """
    Полный текст ячейки: абзацы и вложенные таблицы в порядке OOXML.

    Word часто продолжает большую таблицу как цепочку вложенных <w:tbl> внутри одной ячейки;
    старый код смотрел только на cell.paragraphs и терял основной объём документа.
    """
    parts: list[str] = []
    for block in cell.iter_inner_content():
        if isinstance(block, Paragraph):
            t = _norm(block.text)
            if t:
                parts.append(t)
        elif isinstance(block, Table):
            nested = _nested_table_compact(block)
            if nested:
                parts.append(nested)
    return " | ".join(parts).strip()


def _nested_table_compact(table: Table) -> str:
    """Вложенная таблица внутри ячейки — в одну строку без заголовка [TABLE], чтобы не раздувать markdown."""
    row_chunks: list[str] = []
    for row in table.rows:
        row_chunks.append(" ; ".join(_cell_full_text(c) for c in row.cells))
    return " / ".join(rc for rc in row_chunks if rc.strip())


def _table_to_markdown(table: Table, table_idx: int) -> list[str]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([_cell_full_text(cell) for cell in row.cells])

    if not rows:
        return [f"[TABLE {table_idx}] (empty)", ""]

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    out: list[str] = [f"[TABLE {table_idx}]"]
    header = rows[0]
    out.append("| " + " | ".join(header) + " |")
    out.append("| " + " | ".join(["---"] * width) + " |")
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    out.append("")
    return out


def recognize_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines: list[str] = [f"[DOCX] {path.name}", ""]
    t_idx = 0

    # Порядок блоков как в документе; без смешивания document.paragraphs с body — там ломается индекс
    # из‑за абзацев внутри таблиц.
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = _norm(block.text)
            if text:
                lines.append(text)
                lines.append("")
        elif isinstance(block, Table):
            t_idx += 1
            lines.extend(_table_to_markdown(block, t_idx))

    return "\n".join(lines).strip()
